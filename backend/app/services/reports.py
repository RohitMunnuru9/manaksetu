import json
from io import BytesIO

from docx import Document
from openpyxl import Workbook
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..models import Recommendation, Tender


def record_label(standard) -> str:
    """How a record is named in an exported report.

    A verified record is named by its IS number. An unverified one is named by
    its internal catalogue reference and explicitly says it is not an Indian
    Standard, so an exported file cannot be misread as citing one.
    """
    if standard.standard_number:
        return standard.standard_number
    reference = standard.catalogue_ref or "UNREFERENCED"
    return f"{reference} (internal reference — not an Indian Standard)"


def report_payload(tender: Tender, recommendations: list[Recommendation]) -> dict:
    return {
        "report_type": "ManakSetu recommendation review",
        "tender": {"reference": tender.reference, "title": tender.title, "status": tender.status},
        "safety_notice": "Human review required. Demo or unverified records must not be cited in a tender.",
        "recommendations": [
            {
                "standard_number": item.standard.standard_number,
                "catalogue_ref": item.standard.catalogue_ref,
                "is_an_indian_standard": item.standard.standard_number is not None,
                "title": item.standard.official_title,
                "verification_status": item.standard.verification_status.value,
                "type": item.standard_type,
                "confidence": item.confidence_score,
                "reason": item.reason,
                "matched_requirements": item.matched_requirements,
                "human_review_required": item.human_review_required,
            }
            for item in recommendations
        ],
    }


def build_json(tender: Tender, recommendations: list[Recommendation]) -> bytes:
    return json.dumps(report_payload(tender, recommendations), indent=2, ensure_ascii=False).encode("utf-8")


def build_pdf(tender: Tender, recommendations: list[Recommendation]) -> bytes:
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph("ManakSetu AI", styles["Title"]), Paragraph("Standards Recommendation Review", styles["Heading2"]), Spacer(1, 5 * mm)]
    story += [Paragraph(f"<b>Reference:</b> {tender.reference}", styles["BodyText"]), Paragraph(f"<b>Tender:</b> {tender.title}", styles["BodyText"]), Spacer(1, 4 * mm)]
    warning = Paragraph("HUMAN REVIEW REQUIRED — Demo or unverified records must not be cited in a tender.", styles["BodyText"])
    warning_table = Table([[warning]], colWidths=[170 * mm])
    warning_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HexColor("#FFF3E3")), ("BOX", (0, 0), (-1, -1), 0.5, HexColor("#E8A353")), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [warning_table, Spacer(1, 5 * mm)]
    for index, item in enumerate(recommendations, 1):
        number = record_label(item.standard)
        story += [Paragraph(f"{index}. {number}", styles["Heading3"]), Paragraph(item.standard.official_title, styles["BodyText"]), Paragraph(f"Confidence: {round(item.confidence_score * 100)}% | Status: {item.standard.verification_status.value}", styles["BodyText"]), Paragraph(item.reason, styles["BodyText"]), Spacer(1, 3 * mm)]
    doc.build(story)
    return stream.getvalue()


def build_docx(tender: Tender, recommendations: list[Recommendation]) -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("ManakSetu AI", 0)
    document.add_heading("Standards Recommendation Review", level=1)
    document.add_paragraph(f"Reference: {tender.reference}")
    document.add_paragraph(f"Tender: {tender.title}")
    document.add_paragraph("HUMAN REVIEW REQUIRED — Demo or unverified records must not be cited in a tender.")
    for item in recommendations:
        document.add_heading(record_label(item.standard), level=2)
        document.add_paragraph(item.standard.official_title)
        document.add_paragraph(f"Confidence: {round(item.confidence_score * 100)}% · Verification: {item.standard.verification_status.value}")
        document.add_paragraph(item.reason)
    document.save(stream)
    return stream.getvalue()


def build_xlsx(tender: Tender, recommendations: list[Recommendation]) -> bytes:
    stream = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recommendations"
    sheet.append(["Tender reference", tender.reference])
    sheet.append(["Tender title", tender.title])
    sheet.append(["Safety notice", "Human review required; unverified records must not be cited."])
    sheet.append([])
    sheet.append(["Standard number", "Title", "Type", "Confidence", "Verification", "Human review"])
    for item in recommendations:
        sheet.append([record_label(item.standard), item.standard.official_title, item.standard_type, item.confidence_score, item.standard.verification_status.value, item.human_review_required])
    sheet.freeze_panes = "A6"
    sheet.column_dimensions["A"].width = 46
    sheet.column_dimensions["B"].width = 60
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 14
    sheet.column_dimensions["E"].width = 18
    workbook.save(stream)
    return stream.getvalue()
